# -*- coding: utf-8 -*-
"""Build the gstack Tutorial No.2 lecture decks (EN + ZH) with python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

NAVY = RGBColor(0x06, 0x1F, 0x32)
NAVY2 = RGBColor(0x0D, 0x2E, 0x46)
ACCENT = RGBColor(0x00, 0xA6, 0xC9)
ACCENT2 = RGBColor(0x00, 0xD4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x9A, 0xA8, 0xB4)
DARK = RGBColor(0x1C, 0x27, 0x33)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
MONO = "Consolas"

LATIN = "Arial"
EA = "微软雅黑"

W = Inches(13.333)
H = Inches(7.5)


def sf(run, size=14, bold=False, color=DARK, name=LATIN, ea=EA, mono=False):
    if mono:
        name = MONO
        ea = MONO
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rfonts = rPr.find(qn('a:latin'))
    if rfonts is None:
        rfonts = OxmlElement('a:latin')
        rPr.append(rfonts)
    rfonts.set('typeface', name)
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set('typeface', ea)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return tb, tf


def bg(slide, color):
    r = rect(slide, 0, 0, W, H, fill=color)
    r.shadow.inherit = False
    return r


def footer(slide, text, dark_bg=False):
    _, tf = textbox(slide, Inches(0.6), Inches(6.92), Inches(9.0), Inches(0.3))
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    sf(r, 9, False, GREY if dark_bg else RGBColor(0x7A, 0x86, 0x93))


def pagenum(slide, n, dark_bg=False):
    tb, tf = textbox(slide, Inches(11.9), Inches(6.92), Inches(0.85), Inches(0.3),
                     align=PP_ALIGN.RIGHT)
    r = tf.paragraphs[0].add_run()
    r.text = str(n)
    sf(r, 9, False, GREY if dark_bg else RGBColor(0x7A, 0x86, 0x93))


# ---------------------------------------------------------------- title slide
def slide_title(prs, kicker, title, subtitle, meta):
    s = blank(prs)
    bg(s, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=ACCENT)
    tb, tf = textbox(s, Inches(0.95), Inches(1.55), Inches(11.4), Inches(0.5))
    r = tf.paragraphs[0].add_run()
    r.text = kicker
    sf(r, 15, True, ACCENT2)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT

    tb, tf = textbox(s, Inches(0.95), Inches(2.15), Inches(11.4), Inches(2.0))
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    sf(r, 40, True, WHITE)

    rect(s, Inches(0.95), Inches(4.15), Inches(1.7), Inches(0.055), fill=ACCENT)

    tb, tf = textbox(s, Inches(0.95), Inches(4.5), Inches(11.0), Inches(1.3))
    for i, line in enumerate(subtitle):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        sf(r, 15, False, RGBColor(0xC6, 0xD4, 0xE0))

    tb, tf = textbox(s, Inches(0.95), Inches(6.35), Inches(11.0), Inches(0.4))
    r = tf.paragraphs[0].add_run()
    r.text = meta
    sf(r, 10.5, False, GREY)
    return s


# ------------------------------------------------------------ section divider
def slide_section(prs, number, title, note):
    s = blank(prs)
    bg(s, NAVY2)
    rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=ACCENT)
    tb, tf = textbox(s, Inches(1.0), Inches(2.5), Inches(11.0), Inches(1.0))
    r = tf.paragraphs[0].add_run()
    r.text = number
    sf(r, 60, True, ACCENT)
    tb, tf = textbox(s, Inches(1.0), Inches(3.55), Inches(11.0), Inches(1.2))
    r = tf.paragraphs[0].add_run()
    r.text = title
    sf(r, 30, True, WHITE)
    if note:
        tb, tf = textbox(s, Inches(1.0), Inches(4.85), Inches(10.5), Inches(0.8))
        r = tf.paragraphs[0].add_run()
        r.text = note
        sf(r, 14, False, RGBColor(0xA9, 0xBC, 0xCB))
    return s


# -------------------------------------------------------------- content slide
def _header(s, title, kicker=None):
    if kicker:
        tb, tf = textbox(s, Inches(0.62), Inches(0.42), Inches(12.0), Inches(0.3))
        r = tf.paragraphs[0].add_run()
        r.text = kicker
        sf(r, 11, True, ACCENT)
        ty, th = Inches(0.72), Inches(0.62)
    else:
        ty, th = Inches(0.44), Inches(0.62)
    tb, tf = textbox(s, Inches(0.62), ty, Inches(12.1), th)
    r = tf.paragraphs[0].add_run()
    r.text = title
    sf(r, 27, True, NAVY)
    rect(s, Inches(0.62), Inches(1.42), Inches(1.5), Inches(0.05), fill=ACCENT)


def _note(s, note):
    if not note:
        return
    y = Inches(6.26)
    rect(s, Inches(0.62), y, Inches(12.1), Inches(0.52), fill=LIGHT)
    tb, tf = textbox(s, Inches(0.82), Inches(6.36), Inches(11.7), Inches(0.36),
                     anchor=MSO_ANCHOR.MIDDLE)
    r = tf.paragraphs[0].add_run()
    r.text = note
    sf(r, 11.5, False, NAVY)


def slide_image(prs, title, img_path, note=None, kicker=None,
                max_w=11.2, max_h=4.45):
    """Full-width picture slide below the standard header."""
    s = blank(prs)
    _header(s, title, kicker)
    from PIL import Image
    im = Image.open(img_path)
    ar = im.width / float(im.height)
    w = max_w
    h = w / ar
    if h > max_h:
        h = max_h
        w = h * ar
    s.shapes.add_picture(img_path, Inches((13.333 - w) / 2),
                         Inches(1.62 + (max_h - h) / 2), Inches(w), Inches(h))
    _note(s, note)
    return s


def slide_bullets(prs, title, bullets, note=None, kicker=None, size=15.5):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    tb, tf = textbox(s, Inches(0.7), Inches(1.72), Inches(11.9), Inches(4.4))
    n = len(bullets)
    gap = 16 if n <= 6 else 12
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "▪   "
        sf(r, size, True, ACCENT)
        r2 = p.add_run()
        r2.text = b
        sf(r2, size, False, DARK)
    _note(s, note)
    return s


def slide_checklist(prs, title, items, note=None, kicker=None):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    tb, tf = textbox(s, Inches(0.7), Inches(1.78), Inches(11.9), Inches(4.3))
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(15)
        r = p.add_run()
        r.text = "☐   "
        sf(r, 15.5, True, ACCENT)
        r2 = p.add_run()
        r2.text = b
        sf(r2, 15.5, False, DARK)
    _note(s, note)
    return s


def slide_steps(prs, title, steps, note=None, kicker=None):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    y0 = Inches(1.78)
    for i, st in enumerate(steps):
        y = y0 + Inches(0.86) * i
        rect(s, Inches(0.72), y, Inches(0.42), Inches(0.42), fill=NAVY,
             shape=MSO_SHAPE.OVAL)
        tb, tf = textbox(s, Inches(0.72), y + Inches(0.055), Inches(0.42),
                         Inches(0.34), align=PP_ALIGN.CENTER)
        r = tf.paragraphs[0].add_run()
        r.text = str(i + 1)
        sf(r, 13, True, WHITE)
        tb, tf = textbox(s, Inches(1.34), y + Inches(0.02), Inches(11.3),
                         Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
        r = tf.paragraphs[0].add_run()
        r.text = st
        sf(r, 14, False, DARK)
    _note(s, note)
    return s


def slide_table(prs, title, headers, rows, note=None, kicker=None,
                col_widths=None, first_bold=False, size=12.5, hdr_size=12):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    nrows, ncols = len(rows) + 1, len(headers)
    avail_w = Inches(12.1)
    if col_widths:
        total = sum(col_widths)
        widths = [int(avail_w * w / total) for w in col_widths]
    else:
        widths = [int(avail_w / ncols)] * ncols
    top = Inches(1.72)
    max_h = Inches(4.4)
    row_h = min(int(max_h / nrows), Inches(0.52))
    shp = s.shapes.add_table(nrows, ncols, Inches(0.62), top, avail_w,
                             row_h * nrows)
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    tbl.rows[0].height = Inches(0.42)
    for i in range(1, nrows):
        tbl.rows[i].height = row_h

    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = ""
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = h
        sf(r, hdr_size, True, WHITE)

    for ri, row in enumerate(rows, start=1):
        for c in range(ncols):
            cell = tbl.cell(ri, c)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = row[c] if c < len(row) else ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = val
            sf(r, size, first_bold and c == 0, DARK)
    _note(s, note)
    return s


def slide_code(prs, title, code, note=None, kicker=None, lang_note=None):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    rect(s, Inches(0.62), Inches(1.72), Inches(12.1), Inches(4.15), fill=NAVY)
    tb, tf = textbox(s, Inches(0.95), Inches(1.95), Inches(11.5), Inches(3.7))
    for i, line in enumerate(code):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = line if line else " "
        if line.startswith('#'):
            sf(r, 12, False, RGBColor(0x6E, 0xCF, 0xA8), mono=True)
        else:
            sf(r, 12.5, False, RGBColor(0xE6, 0xF1, 0xF8), mono=True)
    if lang_note:
        tb, tf = textbox(s, Inches(0.95), Inches(5.93), Inches(11.5), Inches(0.25))
        r = tf.paragraphs[0].add_run()
        r.text = lang_note
        sf(r, 10.5, False, GREY)
    _note(s, note)
    return s


def slide_two_col(prs, title, left_title, left_items, right_title, right_items,
                  note=None, kicker=None):
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    for x, t, items in ((Inches(0.62), left_title, left_items),
                        (Inches(6.95), right_title, right_items)):
        rect(s, x, Inches(1.72), Inches(5.78), Inches(0.5), fill=NAVY)
        tb, tf = textbox(s, x + Inches(0.2), Inches(1.79), Inches(5.4),
                         Inches(0.36), anchor=MSO_ANCHOR.MIDDLE)
        r = tf.paragraphs[0].add_run()
        r.text = t
        sf(r, 14, True, WHITE)
        rect(s, x, Inches(2.22), Inches(5.78), Inches(3.85), fill=LIGHT)
        tb, tf = textbox(s, x + Inches(0.25), Inches(2.42), Inches(5.3),
                         Inches(3.5))
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(11)
            r = p.add_run()
            r.text = "▪  "
            sf(r, 13, True, ACCENT)
            r2 = p.add_run()
            r2.text = it
            sf(r2, 13, False, DARK)
    _note(s, note)
    return s


def slide_layers(prs, title, layers, note=None, kicker=None):
    """layers: list of (label, body_lines, color)"""
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    top = Inches(1.75)
    h = Inches(1.35)
    gap = Inches(0.24)
    for i, (label, body, color) in enumerate(layers):
        y = top + (h + gap) * i
        rect(s, Inches(0.62), y, Inches(2.55), h, fill=color)
        tb, tf = textbox(s, Inches(0.78), y + Inches(0.12), Inches(2.3),
                         Inches(1.1), anchor=MSO_ANCHOR.MIDDLE)
        r = tf.paragraphs[0].add_run()
        r.text = label
        sf(r, 14, True, WHITE)
        rect(s, Inches(3.32), y, Inches(9.4), h, fill=LIGHT)
        tb, tf = textbox(s, Inches(3.58), y + Inches(0.14), Inches(9.0),
                         Inches(1.08), anchor=MSO_ANCHOR.MIDDLE)
        for j, line in enumerate(body):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(3)
            r = p.add_run()
            r.text = line
            sf(r, 12.5, j == 0, NAVY)
    _note(s, note)
    return s


def slide_flow(prs, title, boxes, note=None, kicker=None):
    """boxes: list of (x_in, y_in, w_in, h_in, text, color, white_text)"""
    s = blank(prs)
    bg(s, WHITE)
    _header(s, title, kicker)
    for (x, y, w, h, text, color, wt) in boxes:
        rect(s, Inches(x), Inches(y), Inches(w), Inches(h), fill=color)
        tb, tf = textbox(s, Inches(x) + Inches(0.1), Inches(y),
                         Inches(w) - Inches(0.2), Inches(h),
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        r = tf.paragraphs[0].add_run()
        r.text = text
        sf(r, 12, True, WHITE if wt else NAVY)
    _note(s, note)
    return s


def slide_quote(prs, quote, attribution, kicker=None):
    s = blank(prs)
    bg(s, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=ACCENT)
    if kicker:
        tb, tf = textbox(s, Inches(1.0), Inches(1.6), Inches(11.0), Inches(0.4))
        r = tf.paragraphs[0].add_run()
        r.text = kicker
        sf(r, 13, True, ACCENT2)
    tb, tf = textbox(s, Inches(1.0), Inches(2.35), Inches(11.1), Inches(2.6))
    r = tf.paragraphs[0].add_run()
    r.text = quote
    sf(r, 29, True, WHITE)
    rect(s, Inches(1.0), Inches(4.95), Inches(1.7), Inches(0.05), fill=ACCENT)
    tb, tf = textbox(s, Inches(1.0), Inches(5.3), Inches(11.0), Inches(0.8))
    r = tf.paragraphs[0].add_run()
    r.text = attribution
    sf(r, 14, False, RGBColor(0xA9, 0xBC, 0xCB))
    return s


def build(path, slides_fn):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slides_fn(prs)
    prs.save(path)
    print("WROTE", path, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
