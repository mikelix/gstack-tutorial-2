# -*- coding: utf-8 -*-
"""Minimal Markdown -> DOCX renderer tuned for the gstack tutorial deliverables."""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x06, 0x1F, 0x32)
ACCENT = RGBColor(0x00, 0x6E, 0xA6)
GREY = RGBColor(0x55, 0x5F, 0x6D)


def set_font(run, name="Calibri", ea="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), ea)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_rich(par, text, size=10.5, base_bold=False, color=None,
             name="Calibri", ea="微软雅黑"):
    """Handle **bold** and `code` inline."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = par.add_run(part[2:-2])
            set_font(r, name, ea, size, True, color)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = par.add_run(part[1:-1])
            set_font(r, "Consolas", "Consolas", size - 0.5, base_bold,
                     RGBColor(0xA3, 0x1D, 0x1D) if color is None else color)
        else:
            r = par.add_run(part)
            set_font(r, name, ea, size, base_bold, color)


def is_table_sep(line):
    return bool(re.match(r'^\s*\|?[\s:\-|]+\|[\s:\-|]*$', line)) and '-' in line


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def render(md_path, docx_path, title_override=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.3)
        s.right_margin = Cm(2.3)

    i = 0
    n = len(lines)
    first_h1 = True

    def add_code_block(code_lines):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        shade(cell, 'F4F6F8')
        cell.text = ''
        for j, cl in enumerate(code_lines):
            p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(cl if cl else ' ')
            set_font(r, "Consolas", "Consolas", 8.8, False, RGBColor(0x1A, 0x1A, 0x1A))
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code_block(buf)
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'), 'single')
            bot.set(qn('w:sz'), '6')
            bot.set(qn('w:color'), 'C7D0D9')
            pbdr.append(bot)
            pPr.append(pbdr)
            i += 1
            continue

        # table
        if stripped.startswith('|') and i + 1 < n and is_table_sep(lines[i + 1]):
            header = split_row(lines[i])
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            ncol = len(header)
            tbl = doc.add_table(rows=1, cols=ncol)
            tbl.style = 'Table Grid'
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = tbl.rows[0].cells
            for c, h in enumerate(header):
                hdr[c].text = ''
                p = hdr[c].paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                add_rich(p, h, size=9.5, base_bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                         name="Calibri", ea="微软雅黑")
                shade(hdr[c], '061F32')
            for row in rows:
                cells = tbl.add_row().cells
                for c in range(ncol):
                    val = row[c] if c < len(row) else ''
                    val = val.replace('<br>', ' / ')
                    cells[c].text = ''
                    p = cells[c].paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    add_rich(p, val, size=9.2)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # headings
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and first_h1:
                first_h1 = False
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(text)
                set_font(r, "Calibri", "微软雅黑", 24, True, NAVY)
                i += 1
                continue
            sizes = {1: 20, 2: 16, 3: 13, 4: 11.5, 5: 11, 6: 11}
            colors = {1: NAVY, 2: NAVY, 3: ACCENT, 4: NAVY, 5: NAVY, 6: NAVY}
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            set_font(r, "Calibri", "微软雅黑", sizes.get(level, 11), True,
                     colors.get(level, NAVY))
            if level <= 1:
                pPr = p._p.get_or_add_pPr()
                pbdr = OxmlElement('w:pBdr')
                bot = OxmlElement('w:bottom')
                bot.set(qn('w:val'), 'single')
                bot.set(qn('w:sz'), '10')
                bot.set(qn('w:color'), '00A6C9')
                pbdr.append(bot)
                pPr.append(pbdr)
            i += 1
            continue

        # blockquote
        if stripped.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            shade(cell, 'EAF4F9')
            cell.text = ''
            for j, bl in enumerate(buf):
                if not bl:
                    continue
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                add_rich(p, bl, size=10.2, base_bold=False, color=NAVY)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # bullets / checkboxes / numbered
        mb = re.match(r'^[-*]\s+(.*)$', stripped)
        mn = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if mb:
            text = mb.group(1)
            chk = re.match(r'^\[( |x|X)\]\s*(.*)$', text)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            if chk:
                mark = '☑ ' if chk.group(1).lower() == 'x' else '☐ '
                r0 = p.add_run(mark)
                set_font(r0, "Calibri", "微软雅黑", 10.5, False, ACCENT)
                add_rich(p, chk.group(2), size=10.5)
            else:
                add_rich(p, text, size=10.5)
            i += 1
            continue
        if mn:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, mn.group(2), size=10.5)
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        add_rich(p, stripped, size=10.5)
        i += 1

    doc.save(docx_path)
    print('WROTE', docx_path)


if __name__ == '__main__':
    render(sys.argv[1], sys.argv[2])
